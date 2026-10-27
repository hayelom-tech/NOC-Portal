"""Parse the monthly shift roster workbook and its L2 support sheet."""

from __future__ import annotations

import calendar
import csv
import re
import zipfile
from collections.abc import Collection
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

SHIFT_NAME_BY_CODE = {
    "M": "morning",
    "A": "afternoon",
    "N": "night",
}

MAX_XLSX_ENTRIES = 500
MAX_XLSX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_XLSX_ENTRY_BYTES = 32 * 1024 * 1024
MAX_CSV_ROWS = 20_000
MAX_CSV_COLUMNS = 512
MAX_CSV_FIELD_CHARS = 100_000

class RosterWorkbookError(ValueError):
    pass


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _canonical_l2_domain(value: Any, sheet_name: str, row_number: int) -> str:
    normalized = re.sub(r"[^a-z0-9]", "", _cell_text(value).casefold())
    domains = {
        "tx": "tx",
        "transmission": "tx",
        "ran": "ran",
        "radioaccessnetwork": "ran",
        "radioaccessnetworkran": "ran",
        "power": "power",
        "pwr": "power",
        "pscore": "ps_core_sdm",
        "pscoresdm": "ps_core_sdm",
        "cscore": "cs_core_volte",
        "cscorevolte": "cs_core_volte",
    }
    domain = domains.get(normalized)
    if domain is None:
        raise RosterWorkbookError(
            f"Sheet '{sheet_name}' row {row_number}: unknown L2 domain '{_cell_text(value)}'."
        )
    return domain


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for date_format in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, date_format).date()
            except ValueError:
                continue
    return None


def _find_roster_sheet(workbook) -> tuple[str, Any]:
    """Find the FO monthly matrix sheet by its FO Name/Domain headers and dates."""
    candidates: list[tuple[str, Any, int, date]] = []
    for name in workbook.sheetnames:
        sheet = workbook[name]
        first_header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        second_header = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), ())
        first_label = _cell_text(first_header[0] if first_header else "").casefold()
        location_label = _cell_text(first_header[1] if len(first_header) > 1 else "").casefold()
        domain_label = _cell_text(second_header[2] if len(second_header) > 2 else "").casefold()
        if (
            first_label not in {"fo name", "name", "person name"}
            or location_label not in {"location", "site"}
            or domain_label not in {"domain", "team", "department"}
        ):
            continue
        header_dates: list[date] = []
        for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            for cell in row:
                parsed_date = _parse_date(cell)
                if parsed_date:
                    header_dates.append(parsed_date)
        if len(header_dates) < 28:
            continue
        candidates.append((name, sheet, len(header_dates), header_dates[0]))

    if not candidates:
        raise RosterWorkbookError(
            "No Front Office roster sheet found. Expected an FO Name column, a Domain label in row 2, and daily dates in row 1."
        )
    candidates.sort(key=lambda item: item[2], reverse=True)
    name, sheet, _, _ = candidates[0]
    return name, sheet


def parse_roster_workbook(
    content: bytes,
    allowed_team_labels: Collection[str] | None = None,
) -> dict:
    """Parse the FO monthly matrix from an Excel workbook."""
    workbook = _load_checked_workbook(content)
    try:
        sheet_name, sheet = _find_roster_sheet(workbook)
        return _parse_front_office_rows(
            sheet_name,
            list(sheet.iter_rows(values_only=True)),
            _parse_l2_sheet(workbook),
            allowed_team_labels=allowed_team_labels,
        )
    except RosterWorkbookError:
        raise
    except Exception as exc:
        raise RosterWorkbookError("The Excel workbook could not be read. Check that it is a valid, unencrypted .xlsx file.") from exc
    finally:
        workbook.close()


def _load_checked_workbook(content: bytes):
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if len(entries) > MAX_XLSX_ENTRIES:
                raise RosterWorkbookError("The workbook contains too many internal files.")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise RosterWorkbookError("The uploaded file is not a valid .xlsx workbook.")
            expanded_bytes = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise RosterWorkbookError("Encrypted Excel workbooks are not supported.")
                expanded_bytes += entry.file_size
                if entry.file_size > MAX_XLSX_ENTRY_BYTES:
                    raise RosterWorkbookError("A workbook component exceeds the allowed size.")
                if entry.file_size > 1_000_000 and entry.compress_size and entry.file_size / entry.compress_size > 200:
                    raise RosterWorkbookError("The workbook has an unsafe compression ratio.")
            if expanded_bytes > MAX_XLSX_EXPANDED_BYTES:
                raise RosterWorkbookError("The expanded workbook exceeds the 64 MB processing limit.")
    except RosterWorkbookError:
        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise RosterWorkbookError("The uploaded file is not a readable .xlsx workbook.") from exc
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > MAX_CSV_ROWS or (sheet.max_column or 0) > MAX_CSV_COLUMNS:
                workbook.close()
                raise RosterWorkbookError(
                    f"Sheet '{sheet.title}' exceeds the {MAX_CSV_ROWS} row or {MAX_CSV_COLUMNS} column limit."
                )
        return workbook
    except RosterWorkbookError:
        raise
    except Exception as exc:
        raise RosterWorkbookError("The Excel workbook could not be opened. Check that it is a valid .xlsx file.") from exc


def parse_roster_csv(
    content: bytes,
    filename: str = "",
    allowed_team_labels: Collection[str] | None = None,
) -> dict:
    """Parse the FO monthly matrix exported as CSV."""
    rows = _decode_csv_rows(content)
    year = _year_from_filename(filename)
    return _parse_front_office_rows(
        "FO roster CSV",
        rows,
        date_year=year,
        allowed_team_labels=allowed_team_labels,
    )


def _year_from_filename(filename: str) -> int | None:
    match = re.search(r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:tember)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)[ _-]+(20\d{2})", filename, re.IGNORECASE)
    if match:
        return int(match.group(1))
    match = re.search(r"\b(20\d{2})\b", filename)
    return int(match.group(1)) if match else None


def _decode_csv_rows(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise RosterWorkbookError("The CSV file must use UTF-8 or Windows-1252 text encoding.") from exc
    rows: list[list[str]] = []
    try:
        for row_number, row in enumerate(csv.reader(StringIO(text, newline="")), start=1):
            if row_number > MAX_CSV_ROWS:
                raise RosterWorkbookError("The CSV exceeds the 20,000 row limit.")
            if len(row) > MAX_CSV_COLUMNS:
                raise RosterWorkbookError(f"CSV row {row_number} exceeds the 512 column limit.")
            if any(len(value) > MAX_CSV_FIELD_CHARS for value in row):
                raise RosterWorkbookError(f"CSV row {row_number} contains a field that is too large.")
            rows.append(row)
    except csv.Error as exc:
        raise RosterWorkbookError("The CSV file is malformed or contains a field that is too large.") from exc
    return rows


def _parse_front_office_rows(
    sheet_name: str,
    rows: list[tuple | list],
    l2_roster: list[dict] | None = None,
    date_year: int | None = None,
    allowed_team_labels: Collection[str] | None = None,
) -> dict:
    if len(rows) < 3:
        raise RosterWorkbookError("The Front Office roster is missing its headers or assignment rows.")
    first_header = rows[0]
    second_header = rows[1]
    first_label = _cell_text(first_header[0] if first_header else "").casefold()
    location_label = _cell_text(first_header[1] if len(first_header) > 1 else "").casefold()
    domain_label = _cell_text(second_header[2] if len(second_header) > 2 else "").casefold()
    if first_label not in {"fo name", "name", "person name"} or location_label not in {"location", "site"} or domain_label not in {"domain", "team", "department"}:
        raise RosterWorkbookError(
            "This is not an FO roster. Expected FO Name, Location, a Domain label in row 2, and daily date columns."
        )

    date_columns = {}
    for index, cell in enumerate(first_header):
        parsed_date = _parse_date(cell)
        if parsed_date is None and date_year is not None:
            text = _cell_text(cell)
            for pattern in ("%d-%b", "%d-%B", "%d/%b", "%d/%B"):
                try:
                    parsed_date = datetime.strptime(f"{text}-{date_year}", f"{pattern}-%Y").date()
                    break
                except ValueError:
                    continue
        if parsed_date is not None:
            date_columns[index] = parsed_date
    if len(date_columns) < 28:
        date_hint = (
            " For CSV exports with day-month headers (for example, 1-Oct), include the month and year in the filename, "
            "such as October 2026.csv."
            if sheet_name == "FO roster CSV" else ""
        )
        raise RosterWorkbookError(
            f"Sheet '{sheet_name}' must have a full month of daily date columns in row 1.{date_hint}"
        )

    date_values = list(date_columns.values())
    month = date_values[0].month
    year = date_values[0].year
    expected_dates = {
        date(year, month, day)
        for day in range(1, calendar.monthrange(year, month)[1] + 1)
    }
    if set(date_values) != expected_dates:
        raise RosterWorkbookError(
            f"Sheet '{sheet_name}' must contain each date in "
            f"{calendar.month_name[month]} {year} exactly once."
        )

    members: list[dict] = []
    assignments: list[dict] = []
    member_names: set[str] = set()
    normalized_team_labels = (
        {label.strip().casefold() for label in allowed_team_labels}
        if allowed_team_labels is not None
        else None
    )
    email_column = next((
        index
        for header_row in (first_header, second_header)
        for index, value in enumerate(header_row)
        if _normalize_header(value) in {"email", "e_mail"}
    ), None)

    for row_number, row in enumerate(rows[2:], start=3):
        if not row:
            continue
        name = _cell_text(row[0] if len(row) > 0 else None)
        if not name or name.startswith("="):
            continue
        if name in member_names:
            raise RosterWorkbookError(f"Sheet '{sheet_name}' row {row_number}: duplicate member '{name}'.")
        member_names.add(name)

        location = _cell_text(row[1] if len(row) > 1 else None) or None
        team = _cell_text(row[2] if len(row) > 2 else None) or None
        if team and normalized_team_labels is not None and team.casefold() not in normalized_team_labels:
            raise RosterWorkbookError(
                f"Sheet '{sheet_name}' row {row_number}: unknown team label '{team}'. "
                "Add this label in Admin > Team mappings, then upload the roster again."
            )
        email = _cell_text(row[email_column] if email_column is not None and len(row) > email_column else None) or None
        members.append({"name": name, "location": location, "team": team, "email": email})

        for column_index, assignment_date in date_columns.items():
            if column_index >= len(row):
                continue
            shift_code = _cell_text(row[column_index]).upper()
            if not shift_code:
                continue
            if len(shift_code) > 4:
                raise RosterWorkbookError(
                    f"Sheet '{sheet_name}' row {row_number}: shift code must be 4 characters or fewer."
                )
            assignments.append({
                "member_name": name,
                "assignment_date": assignment_date,
                "shift_code": shift_code,
                "shift_name": SHIFT_NAME_BY_CODE.get(shift_code),
            })

    return {
        "month_label": f"{calendar.month_name[month]} {year}",
        "sheet_name": sheet_name,
        "members": members,
        "assignments": assignments,
        "l2_roster": l2_roster or [],
    }


def parse_l2_on_call_workbook(content: bytes) -> dict:
    """Parse the monthly L2 contact roster without per-shift date columns."""
    workbook = _load_checked_workbook(content)
    try:
        selected = None
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            header = {_normalize_header(value) for value in header_row}
            has_domain = bool(header & {"domain", "department"})
            has_name = bool(header & {"name", "engineer_name", "engineername"})
            if has_domain and has_name:
                selected = (sheet.title, [header_row, *rows])
                break
        if selected is None:
            raise RosterWorkbookError(
                "L2 roster needs Domain and Name columns. Tier, Phone, and Email are optional."
            )
        sheet_name, rows = selected
        return _parse_l2_monthly_rows(sheet_name, rows)
    except RosterWorkbookError:
        raise
    except Exception as exc:
        raise RosterWorkbookError("The L2 Excel workbook could not be read. Check that it is a valid .xlsx file.") from exc
    finally:
        workbook.close()


def parse_l2_on_call_csv(content: bytes) -> dict:
    """Parse the monthly L2 contact roster exported as CSV."""
    return _parse_l2_monthly_rows("L2 on-call CSV", _decode_csv_rows(content))


def _normalize_header(value: Any) -> str:
    return "".join(char for char in _cell_text(value).casefold() if char.isalnum() or char == "_")


def _parse_l2_monthly_rows(sheet_name: str, rows: list[tuple | list]) -> dict:
    if not rows:
        raise RosterWorkbookError("The L2 on-call roster is empty.")
    header = [_normalize_header(value) for value in rows[0]]
    columns = {name: index for index, name in enumerate(header) if name}
    domain_column = next((columns[name] for name in ("domain", "department") if name in columns), None)
    name_column = next((columns[name] for name in ("name", "engineer_name", "engineername") if name in columns), None)
    if domain_column is None or name_column is None:
        raise RosterWorkbookError(
            "This is not a monthly L2 roster. Required columns: Domain and Name. Optional columns: Tier, Phone, Email."
        )
    records = []
    current_domain = ""
    for row_number, row in enumerate(rows[1:], start=2):
        if not row or not any(value is not None for value in row):
            continue
        row_domain = _cell_text(row[domain_column] if len(row) > domain_column else "")
        if row_domain:
            current_domain = row_domain
        engineer_name = _cell_text(row[name_column] if len(row) > name_column else "")
        if not engineer_name:
            continue
        if not current_domain:
            raise RosterWorkbookError(f"Sheet '{sheet_name}' row {row_number}: domain is required for each contact group.")

        def optional(name: str) -> str | None:
            aliases = {"tier": ("tier", "role"), "phone": ("phone", "telephone"), "email": ("email", "e_mail")}
            index = next((columns[key] for key in aliases.get(name, (name,)) if key in columns), None)
            value = _cell_text(row[index]) if index is not None and index < len(row) else ""
            return value or None

        records.append({
            "department": _canonical_l2_domain(current_domain, sheet_name, row_number),
            "engineer_name": engineer_name,
            "tier": optional("tier"),
            "week": None,
            "phone": optional("phone"),
            "email": optional("email"),
        })
    if not records:
        raise RosterWorkbookError(f"Sheet '{sheet_name}' contains no L2 contacts.")
    return {"sheet_name": sheet_name, "records": records}


def _parse_l2_sheet(workbook) -> list[dict]:
    """Read the legacy Department/Week sheet or the Domain/Tier contact layout."""
    rows = []
    header = []
    for sheet in workbook.worksheets:
        candidate_rows = list(sheet.iter_rows(values_only=True))
        if not candidate_rows:
            continue
        candidate_header = [_cell_text(cell).casefold() for cell in candidate_rows[0]]
        has_domain = any(value in {"department", "domain"} for value in candidate_header)
        has_name = any(value in {"name", "engineer name", "engineer_name"} for value in candidate_header)
        if has_domain and has_name:
            rows, header = candidate_rows, candidate_header
            break
    if not rows:
        return []

    def column(*names: str) -> int | None:
        return next((i for i, value in enumerate(header) if value in names), None)

    department_column = column("department", "domain")
    week_column = column("week")
    tier_column = column("tier", "role")
    name_column = column("name", "engineer name", "engineer_name")
    phone_column = column("phone")
    email_column = column("email")
    if department_column is None or name_column is None:
        return []

    records: list[dict] = []
    current_department: str | None = None
    required_column = max(department_column, name_column)
    for row_number, row in enumerate(rows[1:], start=2):
        if not row or len(row) <= required_column:
            continue
        department = _cell_text(row[department_column])
        if department:
            current_department = department
        if current_department is None:
            continue

        week = None
        if week_column is not None and len(row) > week_column:
            try:
                week = int(row[week_column])
            except (TypeError, ValueError):
                week = None
        engineer_name = _cell_text(row[name_column])
        if not engineer_name:
            continue
        records.append({
            "department": _canonical_l2_domain(current_department, sheet.title, row_number),
            "week": week,
            "tier": _cell_text(row[tier_column]) if tier_column is not None and len(row) > tier_column else None,
            "engineer_name": engineer_name,
            "phone": _cell_text(row[phone_column]) if phone_column is not None and len(row) > phone_column else None,
            "email": _cell_text(row[email_column]) if email_column is not None and len(row) > email_column else None,
        })
    return records
