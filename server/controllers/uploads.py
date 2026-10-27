"""HTTP routes for the uploads domain."""
import hashlib
import logging
from datetime import date, time
from io import BytesIO
from pathlib import Path
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session, selectinload
from server.config import MAX_UPLOAD_BYTES
from server.controllers.common import record_audit, require_permission
from server.database import get_db
from server.domain.regions import REGION_NAMES
from server.models import ImportLog, L2Roster, RosterAssignment, RosterMember, RosterTeam, Shift, ShiftName, UploadedFile, User, UserRole, utc_now
from server.schemas import ImportLogRead, ImportResult, ImportSheetResult, RosterTeamDomainUpdate, RosterTeamInput, RosterTeamRead
from server.services.permissions import Permission
from server.services.roster_import import RosterWorkbookError, parse_l2_on_call_csv, parse_l2_on_call_workbook, parse_roster_csv, parse_roster_workbook
from server.services.security import get_current_user, require_roles

router = APIRouter()
logger = logging.getLogger(__name__)


@router.get("/api/v1/roster-teams", response_model=list[RosterTeamRead])
def list_roster_teams(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(UserRole.admin)),
):
    return db.scalars(select(RosterTeam).order_by(RosterTeam.name)).all()


@router.post("/api/v1/roster-teams", status_code=201, response_model=RosterTeamRead)
def create_roster_team(
    payload: RosterTeamInput,
    db: Session = Depends(get_db),
    actor: User = Depends(require_roles(UserRole.admin)),
):
    name = payload.name.strip()
    if db.scalar(select(RosterTeam.id).where(func.lower(RosterTeam.name) == name.casefold())):
        raise HTTPException(status_code=409, detail="That roster team label is already registered")
    team = RosterTeam(name=name, challenge_domain=payload.challenge_domain)
    db.add(team)
    try:
        db.flush()
        record_audit(db, actor, "roster_team.created", "roster_team", team.id)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="That roster team label is already registered") from exc
    db.refresh(team)
    return team


@router.patch("/api/v1/roster-teams/{team_id}", response_model=RosterTeamRead)
def update_roster_team_domain(
    team_id: int,
    payload: RosterTeamDomainUpdate,
    db: Session = Depends(get_db),
    actor: User = Depends(require_roles(UserRole.admin)),
):
    team = db.get(RosterTeam, team_id)
    if team is None:
        raise HTTPException(status_code=404, detail="Roster team not found")
    team.challenge_domain = payload.challenge_domain
    record_audit(db, actor, "roster_team.domain_updated", "roster_team", team.id)
    db.commit()
    db.refresh(team)
    return team


@router.post("/api/v1/uploads/excel", status_code=201, response_model=ImportResult)
def upload_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    actor: User = Depends(require_permission(Permission.upload_roster)),
):
    original_name = Path(file.filename or "").name
    extension = Path(original_name).suffix.casefold()
    if extension not in {".xlsx", ".csv"}:
        raise HTTPException(status_code=415, detail="Upload an Excel .xlsx workbook or .csv roster")
    if len(original_name) > 255:
        raise HTTPException(status_code=422, detail="Workbook filename must be 255 characters or fewer")
    content = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Workbook exceeds the 10 MB upload limit")
    if extension == ".xlsx" and not content.startswith(b"PK"):
        raise HTTPException(status_code=415, detail="The uploaded file is not a valid .xlsx workbook")

    checksum = hashlib.sha256(content).hexdigest()
    roster_teams = db.scalars(select(RosterTeam)).all()
    team_ids_by_name = {team.name.casefold(): team.id for team in roster_teams}
    allowed_team_labels = set(team_ids_by_name)
    db.rollback()

    try:
        parsed = (
            parse_roster_csv(content, original_name, allowed_team_labels)
            if extension == ".csv"
            else parse_roster_workbook(content, allowed_team_labels)
        )
    except RosterWorkbookError as exc:
        with db.begin():
            failed_upload = db.scalar(select(UploadedFile).where(UploadedFile.sha256 == checksum))
            if failed_upload and failed_upload.status == "success":
                raise HTTPException(status_code=409, detail="This workbook has already been imported") from exc
            if failed_upload is None:
                failed_upload = UploadedFile(
                    original_name=original_name,
                    sha256=checksum,
                    content=content,
                    uploaded_by_id=actor.id,
                    status="failed",
                )
                db.add(failed_upload)
                db.flush()
            else:
                db.query(ImportLog).filter_by(uploaded_file_id=failed_upload.id).delete()
                failed_upload.original_name = original_name
                failed_upload.content = content
                failed_upload.uploaded_by_id = actor.id
                failed_upload.uploaded_at = utc_now()
                failed_upload.status = "failed"
            db.add(ImportLog(
                uploaded_file_id=failed_upload.id,
                sheet_name="workbook",
                row_count=0,
                imported_count=0,
                status="failed",
                error_message=str(exc),
            ))
            record_audit(db, actor, "roster.import_failed", "uploaded_file", failed_upload.id)
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    shift_times = {
        "morning": (time(6, 0), time(14, 0)),
        "afternoon": (time(14, 0), time(22, 0)),
        "night": (time(22, 0), time(6, 0)),
    }

    assignments = parsed["assignments"]
    members = parsed["members"]
    has_member_emails = any((member.get("email") or "").strip() for member in members)

    all_dates = sorted({assignment["assignment_date"] for assignment in assignments})
    date_min, date_max = (all_dates[0], all_dates[-1]) if all_dates else (None, None)

    try:
        with db.begin():
            upload = db.scalar(select(UploadedFile).where(UploadedFile.sha256 == checksum))
            if upload and upload.status == "success" and not has_member_emails:
                raise HTTPException(status_code=409, detail="This workbook has already been imported")
            if upload is None:
                upload = UploadedFile(
                    original_name=original_name,
                    sha256=checksum,
                    content=content,
                    uploaded_by_id=actor.id,
                    status="success",
                )
                db.add(upload)
                db.flush()
            else:
                db.query(ImportLog).filter_by(uploaded_file_id=upload.id).delete()
                upload.original_name = original_name
                upload.content = content
                upload.uploaded_by_id = actor.id
                upload.uploaded_at = utc_now()
                upload.status = "success"

            member_by_name: dict[str, RosterMember] = {}
            for payload in members:
                existing = db.scalar(select(RosterMember).where(RosterMember.name == payload["name"]))
                email = (payload.get("email") or "").strip().lower() or None
                team = (payload.get("team") or "").strip()
                roster_team_id = team_ids_by_name.get(team.casefold()) if team else None
                matched_user = db.scalar(select(User).where(func.lower(User.email) == email)) if email else None
                if existing is None:
                    existing = RosterMember(
                        name=payload["name"],
                        location=payload["location"],
                        roster_team_id=roster_team_id,
                        email=email,
                        user_id=matched_user.id if matched_user else None,
                    )
                    db.add(existing)
                    db.flush()
                else:
                    if payload["location"] and existing.location != payload["location"]:
                        existing.location = payload["location"]
                    if roster_team_id is not None and existing.roster_team_id != roster_team_id:
                        existing.roster_team_id = roster_team_id
                    if email:
                        existing.email = email
                        existing.user_id = matched_user.id if matched_user else None
                member_by_name[payload["name"]] = existing

            if date_min and date_max:
                db.query(RosterAssignment).filter(
                    RosterAssignment.assignment_date >= date_min,
                    RosterAssignment.assignment_date <= date_max,
                ).delete(synchronize_session=False)
                db.flush()

            assignment_rows: list[RosterAssignment] = []
            for assignment in assignments:
                member = member_by_name.get(assignment["member_name"])
                if member is None:
                    continue
                shift_name_value = None
                if assignment["shift_name"] == "morning":
                    shift_name_value = ShiftName.morning
                elif assignment["shift_name"] == "afternoon":
                    shift_name_value = ShiftName.afternoon
                elif assignment["shift_name"] == "night":
                    shift_name_value = ShiftName.night
                assignment_rows.append(RosterAssignment(
                    member_id=member.id,
                    assignment_date=assignment["assignment_date"],
                    shift_code=assignment["shift_code"],
                    shift_name=shift_name_value,
                ))
            if assignment_rows:
                db.bulk_save_objects(assignment_rows)
                db.flush()

            working_keys: set[tuple[date, str]] = set()
            for assignment in assignments:
                if assignment["shift_name"] in ("morning", "afternoon", "night"):
                    working_keys.add((assignment["assignment_date"], assignment["shift_name"]))

            for assignment_date, shift_name_str in working_keys:
                shift_name_enum = ShiftName(shift_name_str)
                existing_shift = db.scalar(
                    select(Shift).where(
                        Shift.shift_date == assignment_date,
                        Shift.shift_name == shift_name_enum,
                    )
                )
                if existing_shift is None:
                    start, end = shift_times[shift_name_str]
                    db.add(Shift(
                        shift_date=assignment_date,
                        shift_name=shift_name_enum,
                        start_time=start,
                        end_time=end,
                        uploaded_by_id=actor.id,
                    ))
            db.flush()

            month_label = parsed["month_label"]
            sheet_results = [
                ImportSheetResult(
                    sheet_name=f"{month_label} (roster)",
                    row_count=len(members),
                    imported_count=len(members),
                ),
                ImportSheetResult(
                    sheet_name=f"{month_label} (assignments)",
                    row_count=len(assignments),
                    imported_count=len(assignments),
                ),
            ]
            for result in sheet_results:
                db.add(ImportLog(
                    uploaded_file_id=upload.id,
                    sheet_name=result.sheet_name,
                    row_count=result.row_count,
                    imported_count=result.imported_count,
                    status="success",
                    error_message=None,
                ))

            record_audit(db, actor, "roster.imported", "uploaded_file", upload.id)
            file_id = upload.id
    except SQLAlchemyError as exc:
        db.rollback()
        logger.exception("Roster import transaction failed")
        if isinstance(exc, IntegrityError) and db.scalar(
            select(UploadedFile.id).where(UploadedFile.sha256 == checksum)
        ):
            raise HTTPException(status_code=409, detail="This workbook has already been imported") from exc
        raise HTTPException(status_code=500, detail="Import failed; no workbook data was applied") from exc

    return ImportResult(
        file_id=file_id,
        original_name=original_name,
        status="success",
        sheets=sheet_results,
    )


@router.post("/api/v1/uploads/l2-excel", status_code=201)
def upload_l2_on_call_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    actor: User = Depends(require_permission(Permission.upload_roster)),
):
    original_name = Path(file.filename or "").name
    extension = Path(original_name).suffix.casefold()
    if extension not in {".xlsx", ".csv"}:
        raise HTTPException(status_code=415, detail="Upload an Excel .xlsx workbook or .csv roster")
    if len(original_name) > 255:
        raise HTTPException(status_code=422, detail="Workbook filename must be 255 characters or fewer")
    content = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Workbook exceeds the 10 MB upload limit")
    if extension == ".xlsx" and not content.startswith(b"PK"):
        raise HTTPException(status_code=415, detail="The uploaded file is not a valid .xlsx workbook")

    checksum = hashlib.sha256(content).hexdigest()
    try:
        parsed = parse_l2_on_call_csv(content) if extension == ".csv" else parse_l2_on_call_workbook(content)
    except RosterWorkbookError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    db.rollback()
    try:
        with db.begin():
            upload = db.scalar(select(UploadedFile).where(UploadedFile.sha256 == checksum))
            if upload and upload.status == "success":
                raise HTTPException(status_code=409, detail="This workbook has already been imported")
            if upload is None:
                upload = UploadedFile(
                    original_name=original_name,
                    sha256=checksum,
                    content=content,
                    uploaded_by_id=actor.id,
                    status="success",
                )
                db.add(upload)
                db.flush()
            else:
                db.query(ImportLog).filter_by(uploaded_file_id=upload.id).delete()
                upload.original_name = original_name
                upload.content = content
                upload.uploaded_by_id = actor.id
                upload.uploaded_at = utc_now()
                upload.status = "success"

            db.query(L2Roster).delete(synchronize_session=False)
            db.flush()
            for record in parsed["records"]:
                db.add(L2Roster(
                    department=record["department"],
                    week=record["week"],
                    tier=record["tier"],
                    engineer_name=record["engineer_name"],
                    phone=record["phone"],
                    email=record["email"],
                ))
            db.add(ImportLog(
                uploaded_file_id=upload.id,
                sheet_name=parsed["sheet_name"],
                row_count=len(parsed["records"]),
                imported_count=len(parsed["records"]),
                status="success",
            ))
            record_audit(db, actor, "l2_monthly_roster.imported", "uploaded_file", upload.id)
            file_id = upload.id
    except HTTPException:
        raise
    except SQLAlchemyError as exc:
        db.rollback()
        logger.exception("L2 on-call import transaction failed")
        raise HTTPException(status_code=500, detail="L2 monthly roster import failed; no roster data was applied") from exc

    return {
        "file_id": file_id,
        "original_name": original_name,
        "status": "success",
        "imported_count": len(parsed["records"]),
        "sheet_name": parsed["sheet_name"],
    }


@router.get("/api/v1/uploads", response_model=list[ImportLogRead])
def list_import_logs(
    db: Session = Depends(get_db),
    _: User = Depends(require_roles(UserRole.admin)),
):
    return db.scalars(select(ImportLog).order_by(ImportLog.id.desc()).limit(500)).all()


@router.get("/api/v1/roster")
def get_roster(
    date_: date = Query(..., alias="date"),
    db: Session = Depends(get_db),
    actor: User = Depends(get_current_user),
):
    rows = db.scalars(
        select(RosterAssignment)
        .options(
            selectinload(RosterAssignment.member).selectinload(RosterMember.roster_team)
        )
        .where(RosterAssignment.assignment_date == date_)
        .order_by(RosterAssignment.shift_name, RosterAssignment.member_id)
    ).all()
    return [
        {
            "member_name": row.member.name,
            "location": row.member.location,
            "team": row.member.roster_team.name if row.member.roster_team else None,
            "assigned_to_current_user": row.member.user_id == actor.id,
            "shift_code": row.shift_code,
            "shift_name": row.shift_name.value if row.shift_name else None,
        }
        for row in rows
    ]


@router.get("/api/v1/l2-roster")
def get_l2_roster(db: Session = Depends(get_db)):
    rows = db.scalars(
        select(L2Roster).order_by(L2Roster.department, L2Roster.week)
    ).all()
    return [
        {
            "department": row.department,
            "week": row.week,
            "tier": row.tier,
            "engineer_name": row.engineer_name,
            "phone": row.phone,
            "email": row.email,
        }
        for row in rows
    ]


@router.get("/api/v1/uploads/template")
def download_excel_template(_: User = Depends(require_permission(Permission.upload_roster))):
    workbook = Workbook()
    definitions = {
        "FrontOfficeShift": [
            "shift_date", "shift_name", "person_name", "title", "phone", "email", "is_primary",
        ],
        "L2Support": [
            "Domain", "Tier", "Name", "Phone", "Email",
        ],
        "ShiftHandover": [
            "shift_date", "shift_name", "start_time", "end_time", "next_shift_name", "update_for_next_shift",
            "operational_challenges", "active_change_requests", "planned_maintenance",
            "active_field_teams", "major_incidents", "resolved_major_incidents",
            "alarms_pending_resolution", "closed_trouble_tickets", "archived_handover_index",
        ],
        "HandoverRegions": [
            "shift_date", "shift_name", "region_name", "corrective_maintenance_count", "notes",
        ],
    }
    first = True
    for name, headers in definitions.items():
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = name
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if name == "HandoverRegions":
            region_validation = DataValidation(
                type="list",
                formula1=f'"{",".join(REGION_NAMES)}"',
                allow_blank=False,
            )
            region_validation.error = "Select one of the seven official NOC regions."
            region_validation.errorTitle = "Invalid region"
            region_validation.showErrorMessage = True
            sheet.add_data_validation(region_validation)
            region_validation.add("C2:C5001")
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="noc-shift-template.xlsx"'},
    )


@router.get("/api/v1/uploads/l2-template")
def download_l2_on_call_template(_: User = Depends(require_permission(Permission.upload_roster))):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "L2Support"
    sheet.append(["Domain", "Tier", "Name", "Phone", "Email"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="noc-l2-on-call-template.xlsx"'},
    )
